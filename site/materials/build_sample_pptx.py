# -*- coding: utf-8 -*-
"""
產生「上課用ppt.pptx」—— 一個極簡 3 頁投影片，
純粹讓學員在 VS Code 用 Office Viewer 類擴充套件打開來看。

用法：  python3 build_sample_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

INK = RGBColor(0x1E, 0x1E, 0x2E)
STEEL = RGBColor(0x3D, 0x5A, 0x80)
GREY = RGBColor(0x6B, 0x72, 0x80)
PAPER = RGBColor(0xF7, 0xF6, 0xF3)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def slide(title, lines, title_size=40):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER

    bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = STEEL
    bar.line.fill.background()

    tb = s.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(11.5), Inches(1.3))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = "Microsoft JhengHei"

    body = s.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(4.0))
    body.text_frame.word_wrap = True
    for i, line in enumerate(lines):
        p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(22)
        r.font.color.rgb = GREY if line.startswith("（") else INK
        r.font.name = "Microsoft JhengHei"
    return s


slide(
    "上課用範例投影片",
    [
        "這是一個示範用的 .pptx 檔。",
        "",
        "課堂上你會把它放進課程資料夾裡，",
        "然後在 VS Code 裡直接打開來看 —— 不用另外開 PowerPoint。",
        "",
        "（如果打不開，記得先裝 Office Viewer 類擴充套件）",
    ],
)

slide(
    "第 2 頁：VS Code 能做的事",
    [
        "• 寫程式、看程式碼（有顏色標示）",
        "• 一個資料夾就是一個專案",
        "• 裝擴充套件之後，還能看 PPT、Word、PDF",
        "• 內建終端機，可以直接執行 Python",
        "• 內建 Git，幫你記錄每次的修改",
    ],
)

slide(
    "第 3 頁：今天的目標",
    [
        "上完這堂課，你會：",
        "",
        "• 自己裝好 VS Code、把它變成你喜歡的樣子",
        "• 用 AI 幫你寫一個小遊戲，並在終端機跑起來",
        "• 做一張屬於自己的個人名片網頁",
    ],
)

out = Path(__file__).parent / "上課用ppt.pptx"
prs.save(out)
print("saved:", out)
